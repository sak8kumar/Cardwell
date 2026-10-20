#!/usr/bin/env python3
"""
Andrew Cardwell RSI Reversal Screener — F&O universe, BSE prices
=================================================================

Cardwell's "reverse divergence" (reversals) are TREND-CONTINUATION signals:

  POSITIVE REVERSAL (bullish, appears in uptrends)
      RSI makes a LOWER low  while  PRICE makes a HIGHER low.
      Target = Y + (X - W)
        W = price low at first RSI trough
        X = highest high between the two troughs
        Y = price low at second RSI trough

  NEGATIVE REVERSAL (bearish, appears in downtrends)
      RSI makes a HIGHER high while  PRICE makes a LOWER high.
      Target = Y - (W - X)
        W = price high at first RSI peak
        X = lowest low between the two peaks
        Y = price high at second RSI peak

Also reported: Cardwell's RSI range regime
  Bull range ~40-80, Bear range ~20-60. A reversal that agrees with the
  regime is flagged as "Aligned".

Universe : NSE F&O list (fo_mktlots.csv) or your own symbols file
Prices   : Yahoo Finance, BSE tickers (SYMBOL.BO)

Usage
-----
  pip install yfinance pandas numpy openpyxl requests
  python cardwell_screener.py                       # daily, default settings
  python cardwell_screener.py --interval 1wk        # weekly
  python cardwell_screener.py --symbols-file fo.txt # your own list (one symbol per line)
  python cardwell_screener.py --fallback-ns         # use NSE price if a .BO ticker fails

Not investment advice. Validate signals on charts before acting.
"""
import argparse
import io
import sys
import time
from datetime import datetime

import numpy as np
import pandas as pd

NSE_FO_URL = "https://nsearchives.nseindia.com/content/fo/fo_mktlots.csv"
INDEX_SYMBOLS = {"NIFTY", "BANKNIFTY", "FINNIFTY", "MIDCPNIFTY", "NIFTYNXT50",
                 "SENSEX", "BANKEX", "SENSEX50"}
# Add any NSE -> Yahoo-BSE ticker exceptions here if a symbol fails to load
YAHOO_BSE_OVERRIDES = {}


# ----------------------------------------------------------------------------
# Universe
# ----------------------------------------------------------------------------
def load_universe(symbols_file=None):
    if symbols_file:
        with open(symbols_file) as f:
            syms = [s.strip().upper() for s in f if s.strip() and not s.startswith("#")]
        return sorted(set(syms))

    import requests
    headers = {"User-Agent": "Mozilla/5.0", "Accept": "text/csv,*/*",
               "Referer": "https://www.nseindia.com/"}
    try:
        s = requests.Session()
        s.get("https://www.nseindia.com", headers=headers, timeout=10)
        r = s.get(NSE_FO_URL, headers=headers, timeout=15)
        r.raise_for_status()
        df = pd.read_csv(io.StringIO(r.text))
        df.columns = [c.strip().upper() for c in df.columns]
        syms = df["SYMBOL"].astype(str).str.strip().str.upper()
        syms = [s for s in syms if s and s != "SYMBOL" and s not in INDEX_SYMBOLS]
        print(f"Loaded {len(set(syms))} F&O stocks from NSE.")
        return sorted(set(syms))
    except Exception as e:
        sys.exit(f"Could not download the F&O list from NSE ({e}).\n"
                 f"Save the symbols to a text file (one per line) and pass --symbols-file.")


# ----------------------------------------------------------------------------
# Data
# ----------------------------------------------------------------------------
def to_yahoo(sym, suffix):
    return YAHOO_BSE_OVERRIDES.get(sym, sym) + suffix if suffix == ".BO" else sym + suffix


def download(symbols, period, interval, suffix=".BO", chunk=50):
    import yfinance as yf
    out = {}
    for i in range(0, len(symbols), chunk):
        batch = symbols[i:i + chunk]
        tickers = [to_yahoo(s, suffix) for s in batch]
        data = yf.download(tickers, period=period, interval=interval, group_by="ticker",
                           auto_adjust=False, threads=True, progress=False)
        for s, t in zip(batch, tickers):
            try:
                df = data[t] if len(tickers) > 1 else data
                df = df[["Open", "High", "Low", "Close", "Volume"]].dropna()
                if len(df) > 60:
                    out[s] = df
            except Exception:
                pass
        time.sleep(1)
    return out


# ----------------------------------------------------------------------------
# Indicators
# ----------------------------------------------------------------------------
def rsi_wilder(close, n=14):
    d = close.diff()
    gain = d.clip(lower=0)
    loss = -d.clip(upper=0)
    ag = gain.ewm(alpha=1 / n, adjust=False, min_periods=n).mean()
    al = loss.ewm(alpha=1 / n, adjust=False, min_periods=n).mean()
    rs = ag / al.replace(0, np.nan)
    return (100 - 100 / (1 + rs)).fillna(100.0).where(al.notna())


def pivots(series, left, right, kind):
    """Indices of confirmed pivot lows/highs (needs `right` bars after)."""
    v = series.values
    idx = []
    for i in range(left, len(v) - right):
        w = v[i - left:i + right + 1]
        if np.isnan(w).any():
            continue
        target = w.min() if kind == "low" else w.max()
        if v[i] == target and (w == target).sum() == 1:
            idx.append(i)
    return idx


def rsi_regime(rsi, lookback):
    r = rsi.iloc[-lookback:].dropna()
    if r.empty:
        return "N/A"
    lo, hi = r.min(), r.max()
    if lo >= 38 and hi >= 65:
        return "Bull (40-80)"
    if hi <= 62 and lo <= 35:
        return "Bear (20-60)"
    return "Neutral"


# ----------------------------------------------------------------------------
# Signal detection
# ----------------------------------------------------------------------------
def scan(df, a):
    res = []
    df = df.copy()
    df["RSI"] = rsi_wilder(df["Close"], a.rsi)
    df["EMA"] = df["Close"].ewm(span=a.ema, adjust=False).mean()
    n = len(df)
    regime = rsi_regime(df["RSI"], a.regime_lookback)
    last_close = float(df["Close"].iloc[-1])
    last_rsi = float(df["RSI"].iloc[-1])
    fresh_from = n - a.recent - a.pivot_right

    def price_ext(i, col, fn):
        lo, hi = max(0, i - a.price_window), min(n, i + a.price_window + 1)
        return fn(df[col].iloc[lo:hi])

    for kind in ("low", "high"):
        pv = pivots(df["RSI"], a.pivot_left, a.pivot_right, kind)
        if len(pv) < 2:
            continue
        p2 = pv[-1]
        if p2 < fresh_from:
            continue
        best = None
        for p1 in reversed(pv[:-1]):
            gap = p2 - p1
            if gap < a.min_gap:
                continue
            if gap > a.max_gap:
                break
            r1, r2 = df["RSI"].iloc[p1], df["RSI"].iloc[p2]
            if kind == "low":
                W, Y = price_ext(p1, "Low", min), price_ext(p2, "Low", min)
                ok = (r1 - r2) >= a.min_rsi_diff and Y > W * (1 + a.min_price_diff / 100)
            else:
                W, Y = price_ext(p1, "High", max), price_ext(p2, "High", max)
                ok = (r2 - r1) >= a.min_rsi_diff and Y < W * (1 - a.min_price_diff / 100)
            if ok and (best is None or abs(r2 - r1) > best[4]):
                best = (p1, W, Y, gap, abs(r2 - r1), r1, r2)
        if best is None:
            continue

        p1, W, Y, gap, _, r1, r2 = best
        after = df["Close"].iloc[p2 + 1:]
        if kind == "low":
            X = float(df["High"].iloc[p1:p2 + 1].max())
            target = Y + (X - W)
            valid = after.empty or after.min() > Y
            trend_ok = last_close > df["EMA"].iloc[-1]
            aligned = regime.startswith("Bull")
            sig = "Positive Reversal"
            stop = Y
        else:
            X = float(df["Low"].iloc[p1:p2 + 1].min())
            target = Y - (W - X)
            valid = after.empty or after.max() < Y
            trend_ok = last_close < df["EMA"].iloc[-1]
            aligned = regime.startswith("Bear")
            sig = "Negative Reversal"
            stop = Y

        if a.only_valid and not valid:
            continue
        if a.trend_filter and not trend_ok:
            continue

        res.append({
            "Signal": sig,
            "Pivot1 Date": df.index[p1].date(),
            "Pivot2 Date": df.index[p2].date(),
            "Bars Apart": gap,
            "RSI @P1": round(r1, 1),
            "RSI @P2": round(r2, 1),
            "Price W (P1)": round(W, 2),
            "Price X (between)": round(X, 2),
            "Price Y (P2)": round(Y, 2),
            "Last Close": round(last_close, 2),
            "Last RSI": round(last_rsi, 1),
            "Cardwell Target": round(target, 2),
            "Upside/Downside %": round((target / last_close - 1) * 100, 1),
            "Invalidation (Y)": round(stop, 2),
            "Still Valid": "Yes" if valid else "No",
            f"vs EMA{a.ema}": "With trend" if trend_ok else "Against trend",
            "RSI Regime": regime,
            "Regime Aligned": "Yes" if aligned else "No",
            "Bars Since P2": n - 1 - p2,
        })
    return res


# ----------------------------------------------------------------------------
# Output
# ----------------------------------------------------------------------------
def write_excel(results, failed, path, a):
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    df = pd.DataFrame(results)
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        params = pd.DataFrame({
            "Parameter": ["Run time", "Interval", "Period", "RSI length", "Pivot left/right",
                          "Gap between pivots (bars)", "Min RSI diff / price diff %", "Fresh signal window (bars)",
                          "Trend EMA", "Trend filter", "Only valid", "Price source"],
            "Value": [datetime.now().strftime("%Y-%m-%d %H:%M"), a.interval, a.period, a.rsi,
                      f"{a.pivot_left}/{a.pivot_right}", f"{a.min_gap}-{a.max_gap}", f"{a.min_rsi_diff} / {a.min_price_diff}%", a.recent,
                      a.ema, a.trend_filter, a.only_valid,
                      "BSE (.BO)" + (" with NSE fallback" if a.fallback_ns else "")]})
        sheets = {}
        for name, label in [("Positive Reversals", "Positive Reversal"),
                            ("Negative Reversals", "Negative Reversal")]:
            sub = df[df["Signal"] == label] if not df.empty else df
            if not sub.empty:
                sub = sub.sort_values(["Regime Aligned", "Bars Since P2"], ascending=[False, True])
            sub.to_excel(xw, sheet_name=name, index=False)
            sheets[name] = sub
        pd.DataFrame({"Symbol not loaded": failed}).to_excel(xw, sheet_name="Not Loaded", index=False)
        params.to_excel(xw, sheet_name="Parameters", index=False)

        head = PatternFill("solid", fgColor="1F3864")
        green = PatternFill("solid", fgColor="E2EFDA")
        red = PatternFill("solid", fgColor="FCE4D6")
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF")
                c.fill = head
                c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.freeze_panes = "B2"
            for i, col in enumerate(ws.columns, 1):
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
                ws.column_dimensions[get_column_letter(i)].width = min(max(10, width + 2), 22)
            if ws.title in ("Positive Reversals", "Negative Reversals"):
                fill = green if ws.title.startswith("Positive") else red
                hdr = [c.value for c in ws[1]]
                if "Regime Aligned" in hdr:
                    k = hdr.index("Regime Aligned")
                    for row in ws.iter_rows(min_row=2):
                        if row[k].value == "Yes":
                            for c in row:
                                c.fill = fill


def main():
    p = argparse.ArgumentParser(description="Cardwell RSI reversal screener (F&O, BSE prices)")
    p.add_argument("--symbols-file")
    p.add_argument("--interval", default="1d", choices=["1d", "1wk"])
    p.add_argument("--period", default="2y")
    p.add_argument("--rsi", type=int, default=14)
    p.add_argument("--pivot-left", type=int, default=5)
    p.add_argument("--pivot-right", type=int, default=2)
    p.add_argument("--price-window", type=int, default=2,
                   help="bars either side of RSI pivot to find the price extreme")
    p.add_argument("--min-gap", type=int, default=5)
    p.add_argument("--max-gap", type=int, default=60)
    p.add_argument("--min-rsi-diff", type=float, default=2.0,
                   help="minimum RSI points between the two pivots")
    p.add_argument("--min-price-diff", type=float, default=0.5,
                   help="minimum %% price difference between the two pivots")
    p.add_argument("--recent", type=int, default=10,
                   help="second pivot must be within this many bars of today")
    p.add_argument("--ema", type=int, default=50)
    p.add_argument("--regime-lookback", type=int, default=60)
    p.add_argument("--trend-filter", action="store_true",
                   help="keep only signals on the right side of the EMA")
    p.add_argument("--only-valid", action="store_true",
                   help="drop signals where price has already closed beyond Y")
    p.add_argument("--fallback-ns", action="store_true")
    p.add_argument("--out", default=None)
    a = p.parse_args()

    syms = load_universe(a.symbols_file)
    print(f"Downloading {len(syms)} symbols from BSE ({a.interval}, {a.period})...")
    data = download(syms, a.period, a.interval, ".BO")
    missing = [s for s in syms if s not in data]
    if missing and a.fallback_ns:
        print(f"{len(missing)} failed on BSE, retrying on NSE...")
        data.update(download(missing, a.period, a.interval, ".NS"))
        missing = [s for s in syms if s not in data]

    results = []
    for s, df in data.items():
        for r in scan(df, a):
            results.append({"Symbol": s, **r})

    out = a.out or f"cardwell_reversals_{a.interval}_{datetime.now():%Y%m%d}.xlsx"
    write_excel(results, missing, out, a)

    pos = sum(r["Signal"] == "Positive Reversal" for r in results)
    neg = len(results) - pos
    print(f"\nScanned {len(data)} | Positive: {pos} | Negative: {neg} | Not loaded: {len(missing)}")
    if results:
        cols = ["Symbol", "Signal", "Pivot2 Date", "Last Close", "Cardwell Target",
                "Upside/Downside %", "RSI Regime", "Regime Aligned"]
        print(pd.DataFrame(results)[cols].to_string(index=False))
    print(f"\nSaved: {out}")


if __name__ == "__main__":
    main()
